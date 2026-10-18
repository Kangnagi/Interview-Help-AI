"""
질문 모델 q3 학습 데이터 검토표 — Claude가 쓴 경력자 자기소개서 · 질문 6개(data/qgen_exp_intro_b*.jsonl)를 엑셀로 만든다.
학습 전에 팀이 질문마다 O/X를 고르고, X가 많은 유형은 남은 데이터를 쓰기 전에 기준을 고친다.

  python training/make_qgen_exp_review.py [배치 파일 …]   (없으면 data/qgen_exp_intro_b*.jsonl 전부)
출력: 바탕화면 '내일의 면접' 폴더에 질문생성검토_q3.xlsx
참고 질문(refs)은 AI Hub '채용면접 인터뷰 데이터'의 실제 면접관 질문 — 데이터 이용 조건상 팀 내부에서만 쓴다.
"""
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OCC = {"BM": "경영·사무", "SM": "영업·마케팅", "ICT": "ICT", "PS": "공공·서비스", "MM": "생산·제조", "ARD": "디자인", "RND": "연구개발"}
JOB_SLOTS = (2, 5)   # 질문 6개 순서: 자기소개서, 자기소개서, 직무, 자기소개서, 자기소개서, 직무 (llama_service.qgen_messages)

paths = sys.argv[1:] or sorted(glob.glob(os.path.join(HERE, "data", "qgen_exp_intro_b*.jsonl")))
items = [json.loads(l) for p in paths for l in open(p, encoding="utf-8") if l.strip()]

wb = Workbook()
ws = wb.active
ws.title = "검토"
head = ["번호", "직군", "지원 정보", "자기소개서", "순서", "질문 종류", "질문", "응용한 실제 면접 질문", "좋은 질문 (O/X)", "메모"]
ws.append(head)
for c in ws[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4F6EF7")
shade = PatternFill("solid", fgColor="F3F5FF")
row = 2
for n, it in enumerate(items, 1):
    info = "\n".join(x for x in (it.get("company") and f"회사: {it['company']}", f"직무: {it['job']}",
                                 f"업무: {it['duty']}", it.get("ideal") and f"인재상: {it['ideal']}") if x)
    intro = it.get("intro") or "(자기소개서 없음 — 6개 모두 직무 질문)"
    refs = "\n".join(f"• {r}" for r in it.get("refs", []))
    start = row
    first_job = JOB_SLOTS[0] if it.get("intro") else 0   # 참고 질문은 첫 직무 질문 줄에만 (같은 칸 반복 방지)
    for i, q in enumerate(it["questions"]):
        kind = "직무" if (i in JOB_SLOTS or not it.get("intro")) else "자기소개서"
        ws.append([it["id"], OCC.get(it["occupation"], it["occupation"]), info if i == 0 else "", intro if i == 0 else "",
                   i + 1, kind, q, refs if i == first_job else "", "", ""])
        row += 1
    for col in "ABCD":
        ws.merge_cells(f"{col}{start}:{col}{row - 1}")
    if n % 2 == 0:
        for r in ws.iter_rows(min_row=start, max_row=row - 1):
            for c in r:
                c.fill = shade
for col, width in zip("ABCDEFGHIJ", (8, 11, 30, 60, 6, 10, 55, 45, 12, 28)):
    ws.column_dimensions[col].width = width
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Arial", size=10)
dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"I2:I{row - 1}")
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:J{row - 1}"

help_ws = wb.create_sheet("안내", 0)
for line in (
    "질문 모델 q3 학습 데이터 검토 (경력자 자기소개서 기반 질문 생성)",
    "",
    f"지원 정보 {len(items)}개 × 질문 6개 = {len(items) * 6}개. '검토' 시트의 I열에 O / X를 고르고, X면 J열에 이유를 짧게 적어 주세요.",
    "직군별로 나눠 검토해도 됩니다 (B열 필터).",
    "",
    "이번 학습 방향 (10/14)",
    "• 꼬리 질문은 새로 학습하지 않고, 면접 시작 전 질문(자기소개서 질문 · 직무 질문)만 학습합니다.",
    "• 실제 면접은 꼬리 질문보다 자기소개서 질문이 많으므로, 질문 6개 중 4개가 자기소개서 질문입니다.",
    "• 직무 질문은 AI Hub 실제 면접관 질문(H열)을 지원 직무 · 경력에 맞게 응용해 썼습니다.",
    "",
    "좋은 질문(O)의 기준 — training/question_guide.md",
    "• 자기소개서에 실제로 쓴 경험 · 수치 · 판단을 짚는다 (없는 내용을 지어내지 않음)",
    "• 한 번에 하나만 묻는다",
    "• 자기소개서에 빠진 것(본인 역할, 판단 근거, 결과 확인 방법, 어려움과 대처)을 묻는다 — 쓰인 걸 다시 말하게 하지 않음",
    "• 실제 면접관이 경력자에게 물을 만한 자연스러운 말투",
    "• 직무 질문: 지원 직무의 실제 업무 상황에서 어떻게 할지 묻는다 (실제 면접 질문을 그대로 베끼지 않고 응용)",
    "",
    "AI Hub 데이터의 실제 면접 질문이 들어 있으니 팀 밖으로 공유하지 마세요.",
):
    help_ws.append([line])
help_ws.column_dimensions["A"].width = 120
help_ws["A1"].font = Font(name="Arial", bold=True, size=14)

out = os.path.join(os.path.expanduser("~"), "Desktop", "내일의 면접", "질문생성검토_q3.xlsx")
wb.save(out)
print(f"{len(items)}개 지원 정보, {len(items) * 6}개 질문 → {out}")
