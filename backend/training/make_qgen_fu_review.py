"""
꼬리 질문 유형 검토표 — eval_qgen.py 결과에서 어댑터가 쓴 꼬리 질문만 모아 엑셀로 만든다 (팀이 유형 · 품질을 O/X로 확인).

  python training/make_qgen_fu_review.py training/data/qgen_eval_bllossom-qgen-q2.jsonl
출력: 바탕화면 '내일의 면접' 폴더에 꼬리질문검토_<이름>.xlsx
실제 면접 답변(바탕화면 원본)이 들어 있으므로 팀 내부에서만 쓰고 외부에 공유하지 않는다.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TYPE_NAMES = {   # services/llm/llama_service.FU_TYPES와 같은 이름
    "role": "경험 진위 및 역할 검증형",
    "why": "판단 근거 및 원인 분석형",
    "whatif": "가정 및 바뀐 조건 대응형",
    "strength": "강점 및 약점 심화 탐색형",
}

src = sys.argv[1]
name = os.path.basename(src).replace("qgen_eval_", "").replace(".jsonl", "")
rows = [r for r in (json.loads(l) for l in open(src, encoding="utf-8")) if r["task"] == "followup"]

wb = Workbook()
ws = wb.active
ws.title = "검토"
ws.append(["번호", "면접 질문", "지원자 답변", "요청 유형", "꼬리 질문 (어댑터)", "유형 맞음 (O/X)", "좋은 질문 (O/X)", "메모"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4F6EF7")
for i, r in enumerate(rows, 1):
    ws.append([i, r["question"], r["answer"], TYPE_NAMES[r["fu_type"]], r["model"] or "(빈 출력)", "", "", ""])
for col, width in zip("ABCDEFGH", (6, 40, 70, 22, 60, 14, 14, 30)):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F2:G{len(rows) + 1}")
ws.freeze_panes = "E2"

help_ws = wb.create_sheet("안내", 0)
for line in (
    "꼬리 질문 유형 검토",
    "",
    "'검토' 시트의 항목마다 어댑터가 쓴 꼬리 질문을 보고 F열(요청한 유형으로 썼는지)과 G열(면접관이 실제로 물을 만한 좋은 질문인지)에 O / X를 고르세요.",
    "같은 답변이 두 번 나오는 것은 1순위 · 2순위 유형으로 각각 쓴 것입니다 (면접 하나에 꼬리 질문이 둘이면 서로 다른 유형으로 묻습니다).",
    "",
    "유형 기준 (10/7 팀 기준)",
    "• 경험 진위 및 역할 검증형: 말한 성과 · 경험에서 본인이 실제로 맡은 역할과 기여도를 확인",
    "• 판단 근거 및 원인 분석형: 그 방법을 고른 이유, 원인을 어떻게 판단했는지",
    "• 가정 및 바뀐 조건 대응형: 조건을 바꿔 '만약 ~라면 어떻게' 대처할지",
    "• 강점 및 약점 심화 탐색형: 말한 강점이 직무에서 어떻게 쓰이는지, 약점이 업무에 지장을 주지 않는지",
    "",
    "실제 면접 답변이 들어 있으니 팀 밖으로 공유하지 마세요.",
):
    help_ws.append([line])
help_ws.column_dimensions["A"].width = 120
help_ws["A1"].font = Font(bold=True, size=14)

out = os.path.join(os.path.expanduser("~"), "Desktop", "내일의 면접", f"꼬리질문검토_{name}.xlsx")
wb.save(out)
print(f"{len(rows)}개 항목 → {out}")
