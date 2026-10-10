"""
팀 블라인드 비교표 — eval_qgen.py 결과로 '지금 방식(질문 틀)'과 '질문 전용 어댑터'의 질문을 A/B로 섞어 엑셀로 만든다.

  python training/make_qgen_blind.py training/data/qgen_eval_bllossom-qgen-q1.jsonl
출력: 바탕화면 '내일의 면접' 폴더에 질문비교_블라인드_<이름>.xlsx
  - '평가' 시트: 항목마다 A안 · B안 질문, 평가자가 더 나은 쪽(A/B/같음)과 메모를 적는다
  - '정답' 시트(숨김): 어느 쪽이 어댑터인지 — 평가가 끝난 뒤 열어 집계
실제 면접 답변(바탕화면 원본)이 들어 있으므로 팀 내부에서만 쓰고 외부에 공유하지 않는다.
"""
import json
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

src = sys.argv[1]
name = os.path.basename(src).replace("qgen_eval_", "").replace(".jsonl", "")
rows = [json.loads(l) for l in open(src, encoding="utf-8")]
rng = random.Random(7)

wb = Workbook()
ws = wb.active
ws.title = "평가"
key = wb.create_sheet("정답")
ws.append(["번호", "종류", "입력 (자기소개서 · 면접 질문과 답변)", "A안", "B안", "더 나은 쪽 (A/B/같음)", "메모"])
key.append(["번호", "A안", "B안"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4F6EF7")

for i, r in enumerate(rows, 1):
    if r["task"] == "qgen":
        kind = "자기소개서 질문"
        src_text = r["resume_text"]
        a, b = "\n".join(f"- {q}" for q in r["template"]), "\n".join(f"- {q}" for q in r["model"])
    else:
        kind = "꼬리 질문"
        src_text = f"[질문] {r['question']}\n[답변] {r['answer']}"
        a, b = r["template"] or "(꼬리 질문 없음)", r["model"] or "(꼬리 질문 없음)"
    labels = ["질문 틀", "어댑터"]
    if rng.random() < 0.5:
        a, b = b, a; labels.reverse()
    ws.append([i, kind, src_text, a, b, "", ""])
    key.append([i, labels[0], labels[1]])

for col, width in zip("ABCDEFG", (6, 14, 70, 60, 60, 16, 30)):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
dv = DataValidation(type="list", formula1='"A,B,같음"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"F2:F{len(rows) + 1}")
ws.freeze_panes = "C2"
key.sheet_state = "hidden"

ws_help = wb.create_sheet("안내", 0)
for line in (
    "질문 비교 블라인드 평가",
    "",
    "'평가' 시트의 항목마다 A안과 B안 중 더 좋은 면접 질문을 고르세요 (F열: A / B / 같음).",
    "기준 (training/question_guide.md 요약): 자기소개서·답변 내용에 근거하는가, 질문이 하나인가, 빠진 부분을 파고드는가,",
    "다른 지원자에게는 할 수 없는 구체적인 질문인가, 자연스러운 면접관 말투인가.",
    "어느 쪽이 어떤 방식인지는 평가가 끝날 때까지 보지 마세요 ('정답' 시트는 숨겨져 있습니다).",
    "실제 면접 답변이 들어 있으니 팀 밖으로 공유하지 마세요.",
):
    ws_help.append([line])
ws_help.column_dimensions["A"].width = 120
ws_help["A1"].font = Font(bold=True, size=14)

out = os.path.join(os.path.expanduser("~"), "Desktop", "내일의 면접", f"질문비교_블라인드_{name}.xlsx")
wb.save(out)
print(f"{len(rows)}개 항목 → {out}")
